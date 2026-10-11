from docx import Document
from typing import Dict
from app.core import logger


class DOCXParser:
    """Parser for DOCX documents."""
    
    def parse(self, file_path: str) -> Dict:
        """
        Parse DOCX document.
        
        Returns:
            Dict with text_content and metadata
        """
        try:
            logger.info(f"Parsing DOCX: {file_path}")
            
            doc = Document(file_path)
            
            # Extract text from paragraphs
            text_content = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)
            
            # Extract text from tables
            table_content = []
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        table_content.append(row_text)
            
            # Combine all text
            full_text = "\n".join(text_content)
            if table_content:
                full_text += "\n\n--- Tables ---\n" + "\n".join(table_content)
            
            # Metadata
            metadata = {
                "total_paragraphs": len(doc.paragraphs),
                "total_tables": len(doc.tables),
                "char_count": len(full_text)
            }
            
            # Try to get core properties
            try:
                core_props = doc.core_properties
                metadata["title"] = core_props.title or ""
                metadata["author"] = core_props.author or ""
                metadata["subject"] = core_props.subject or ""
                metadata["created"] = str(core_props.created) if core_props.created else ""
                metadata["modified"] = str(core_props.modified) if core_props.modified else ""
            except Exception:
                pass
            
            logger.info(f"DOCX parsed successfully: {len(doc.paragraphs)} paragraphs, {len(full_text)} characters")
            
            return {
                "text_content": full_text,
                "metadata": metadata
            }
            
        except Exception as e:
            logger.error(f"Failed to parse DOCX {file_path}: {e}")
            raise


# Global DOCX parser instance
docx_parser = DOCXParser()
